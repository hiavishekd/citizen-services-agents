"""Convert selected sheets of the grounding workbook into CSV files under data/.

Each sheet starts with a one-line note, then a blank row, then the header row.
The header row is the first row with at least three filled cells; the note is dropped.

Usage: python scripts/xlsx_to_csv.py [Sheet ...]   (defaults to the core sheet list)
Requires: openpyxl
"""
import csv
import sys
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
WORKBOOK = DATA_DIR / "civic_platform_grounding_data.xlsx"
DEFAULT_SHEETS = [
    "Sources_Registry", "Helplines", "Schemes", "Personas_Tests",
    "Health_RedFlags", "Fraud_Patterns", "Fraud_Report_Steps", "Gov_Events",
]


def _evaluate(value, ws, row: int):
    """Cached values are lost when the workbook is re-saved by openpyxl, so evaluate the one formula we use."""
    if isinstance(value, str) and value.startswith("=IF(F") and value.endswith('"suppressed")'):
        return ws[f"F{row}"].value if ws[f"F{row}"].value >= ws["B1"].value else "suppressed"
    return value


def convert(sheet_name: str, wb) -> int:
    ws = wb[sheet_name]
    rows = [["" if c.value is None else str(_evaluate(c.value, ws, c.row)).strip() for c in r]
            for r in ws.iter_rows()]
    start = next(i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 3)
    body = [r for r in rows[start:] if any(r)]
    with open(DATA_DIR / f"{sheet_name}.csv", "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(body)
    return len(body) - 1


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK)
    for name in sys.argv[1:] or DEFAULT_SHEETS:
        print(f"{name}.csv: {convert(name, wb)} data rows")


if __name__ == "__main__":
    main()
